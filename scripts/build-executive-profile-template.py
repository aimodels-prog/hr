"""Build a blank, confidential profile-collection workbook (not an automatic importer)."""
from pathlib import Path
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts" / "VIA_Executive_Profile_Template.xlsx"
NAVY, BLUE, TEAL = "073450", "075B98", "008F83"
PALE, INK, GREY = "EAF4FA", "18374C", "5D7180"
wb = Workbook()
start = wb.active
start.title = "Start Here"
lists = wb.create_sheet("Choices")
choices = {
    "YesNo": ["Yes", "No", "Not confirmed"],
    "Status": ["Active", "Onboarding", "Probation", "Notice", "Inactive"],
    "EntryType": ["Existing Employee", "New Employee"],
    "Reporting": ["No supervisor - CEO", "Has supervisor", "HR to confirm"],
    "Gender": ["Male", "Female"],
    "Marital": ["Single", "Married", "Divorced", "Widowed"],
    "Relationship": ["Spouse", "Child", "Parent", "Other"],
    "Visa": ["Yes", "No", "Not confirmed"],
    "Payment": ["Bank Transfer", "Cheque", "Cash"],
    "Frequency": ["Monthly", "Biweekly", "Weekly"],
    "DocumentFor": ["Employee", "Dependant"],
    "Availability": ["Attached separately", "Already in app", "Not available yet", "Not applicable"],
    "Review": ["Not reviewed", "Ready for import review", "Needs clarification"],
    "DocumentType": ["Profile Photo", "Employment Contract", "Passport", "Visa", "Resident Card / National ID", "Work Permit", "Driving Licence", "Medical Document", "Education Degree", "Professional Certificate", "Oman Engineering Certificate", "Bank Evidence", "Health Insurance Card", "Insurance Table of Benefits", "Updated CV", "Other Document"],
}
for col, (key, values) in enumerate(choices.items(), 1):
    lists.cell(1, col, key)
    for row, value in enumerate(values, 2):
        lists.cell(row, col, value)
    letter = lists.cell(1, col).column_letter
    wb.defined_names.add(DefinedName(key, attr_text=f"'Choices'!${letter}$2:${letter}${len(values)+1}"))
lists.sheet_state = "hidden"


def base(ws, title, subtitle, count):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = BLUE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=count)
    ws.cell(1, 1, title).font = Font(name="Calibri", size=21, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 40
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=count)
    ws.cell(2, 1, subtitle).font = Font(name="Calibri", size=11, color=GREY)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 34
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=count)
    ws.cell(3, 1, "CONFIDENTIAL  |  One person can appear on several sheets. Use the same work email each time.")
    ws.cell(3, 1).font = Font(name="Calibri", size=10, color=TEAL, bold=True)
    ws.row_dimensions[3].height = 24
    ws.cell(4, 1, "← Start Here").hyperlink = "#'Start Here'!A1"
    ws.cell(4, 1).font = Font(name="Calibri", color=BLUE, underline="single")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:6"
    ws.oddFooter.center.text = "VIA HR Application | Confidential | Page &P of &N"


def sheet(name, title, note, columns, required=(), dropdowns=None, dates=(), numeric=(), rows=30):
    ws = wb.create_sheet(name)
    base(ws, title, note, len(columns))
    ws.freeze_panes = "C7"
    for col, (heading, hint, width) in enumerate(columns, 1):
        cell = ws.cell(6, col, heading)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL if heading in required else BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.comment = Comment(hint, "VIA HR Application")
        ws.column_dimensions[cell.column_letter].width = width
        for row in range(7, rows + 7):
            c = ws.cell(row, col)
            c.font = Font(name="Calibri", size=11, color=INK)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.number_format = "@"  # IDs, bank numbers and phone numbers must keep leading zeros.
            c.fill = PatternFill("solid", fgColor="F1F8FC" if row % 2 else "FFFFFF")
            ws.row_dimensions[row].height = 30
        target = f"{cell.column_letter}7:{cell.column_letter}{rows+6}"
        if heading in required:
            ws.conditional_formatting.add(target, FormulaRule(
                formula=[f'AND($A7<>"",{cell.column_letter}7="")'],
                fill=PatternFill("solid", fgColor="FFF0D6")))
        if dropdowns and heading in dropdowns:
            dv = DataValidation(type="list", formula1="=" + dropdowns[heading], allow_blank=True)
            dv.showErrorMessage = True
            dv.error = "Please choose a value from the list. Leave blank if not known."
            dv.errorTitle = "Choose an option"
            dv.errorStyle = "stop"
            ws.add_data_validation(dv)
            dv.add(target)
        if heading in dates:
            dv = DataValidation(type="date", operator="between", formula1="DATE(1900,1,1)", formula2="DATE(2200,12,31)", allow_blank=True)
            dv.showErrorMessage = True
            dv.error = "Enter a real date, for example 07-Oct-2026."
            ws.add_data_validation(dv)
            dv.add(target)
            for row in range(7, rows + 7): ws.cell(row, col).number_format = "dd-mmm-yyyy"
        if heading in numeric:
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=True)
            dv.showErrorMessage = True
            dv.error = "Enter zero or a positive number, without a currency symbol."
            ws.add_data_validation(dv)
            dv.add(target)
            for row in range(7, rows + 7): ws.cell(row, col).number_format = "0.00"
    ws.row_dimensions[6].height = 44
    last = ws.cell(6, len(columns)).column_letter
    table = Table(displayName="Data" + name.replace(" ", ""), ref=f"A6:{last}{rows+6}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{last}{rows+6}"
    return ws


EMAIL = ("Work email", "Use the exact VIA email used in the app or portal. This links the person's records across sheets. Do not use an assistant's email.", 30)
NAME = ("Full name", "For checking only; use the same name as the People sheet.", 26)
sheet("People", "01 / PEOPLE", "Start here. One row per person. Green headings are needed for a usable profile; amber cells show missing information once work email is entered.", [
    EMAIL, ("Employee number", "Copy the existing employee number from the app. Do not invent one; if unknown leave blank for matching review.", 22),
    ("Legal name", "Full name exactly as shown on passport or ID.", 30),
    ("Preferred name", "Name the person wants colleagues to use.", 23),
    ("Phone", "Include country code. Store as text, for example +968 followed by the number.", 23),
    ("Personal email", "Optional personal email; do not substitute for work email.", 30),
    ("Already in app?", "Yes if there is already a profile, even if incomplete. Unknown: Not confirmed.", 22),
    ("Prepared by", "Name of HR or executive assistant who collected these details.", 26),
    ("Prepared on", "Date this row was prepared.", 19),
    ("Employee confirmed details?", "Record whether the employee has checked the supplied details. This is not HR approval.", 26),
    ("HR review", "Collection review only; does not approve documents or employment in the app.", 26),
    ("HR reviewer", "Name of the HR person reviewing this collection.", 25),
    ("Notes for HR", "Missing information or questions. Do not enter passwords.", 38),
], required=("Work email", "Legal name", "Preferred name"), dropdowns={"Already in app?":"YesNo", "Employee confirmed details?":"YesNo", "HR review":"Review"}, dates=("Prepared on",))

employment = sheet("Employment", "02 / EMPLOYMENT — HR COMPLETES", "CEO: choose 'No supervisor - CEO' and leave both supervisor fields blank. Country Managers keep their actual supervisor; never use a dummy person.", [
    EMAIL, NAME, ("Position", "Use the exact position name in the app, e.g. CEO or Country Manager.", 27),
    ("Department", "Exact department name in Company Setup.", 25),
    ("Work location", "Exact location name in Company Setup.", 25),
    ("Country", "Country of employment.", 21),
    ("Legal entity", "Employing company / legal entity, if applicable.", 26),
    ("Employment type", "Exact employment type from Company Setup.", 23),
    ("Joining date", "Original VIA joining date, not the date someone logged in.", 19),
    ("Probation end date", "Leave blank if not applicable or not confirmed.", 22),
    ("Staff entry type", "Existing Employee or New Employee.", 24),
    ("Employment status", "Current employment status, confirmed by HR.", 23),
    ("Reporting arrangement", "CEO has no supervisor. This does not grant self-approval rights.", 29),
    ("Supervisor work email", "Leave blank for CEO. For others use the actual supervisor's work email.", 30),
    ("Supervisor employee number", "Existing app employee number, if known. Never guess.", 28),
    ("Visa required?", "HR confirms whether the employee needs a visa.", 22),
    ("Grade", "Exact grade from Company Setup, if used.", 19),
    ("Project", "Exact project name, if assigned.", 28),
    ("Cost centre", "Exact cost centre name, if assigned.", 25),
    ("Weekly working hours", "Contractual working hours for payroll/proration, if applicable. Do not assume a value.", 24),
], required=("Work email", "Position", "Department", "Work location", "Employment type", "Joining date", "Reporting arrangement"), dropdowns={"Staff entry type":"EntryType", "Employment status":"Status", "Reporting arrangement":"Reporting", "Visa required?":"YesNo"}, dates=("Joining date", "Probation end date"), numeric=("Weekly working hours",))
employment.conditional_formatting.add("N7:O36", FormulaRule(formula=['AND(OR(UPPER(TRIM($C7))="CEO",$M7="No supervisor - CEO"),OR($N7<>"",$O7<>""))'], fill=PatternFill("solid", fgColor="FFD9D9")))

sheet("Personal", "03 / PERSONAL DETAILS", "HR or the assistant may collect these from the employee. Leave unknown information blank and explain it in People → Notes for HR.", [
    EMAIL, NAME, ("Date of birth", "Date as shown on the employee's identity document.", 20),
    ("Gender", "Choose the value supported by the current app; leave blank if not supplied.", 17),
    ("Nationality", "Nationality shown on passport / ID.", 24),
    ("Marital status", "Choose only if confirmed by the employee.", 20),
    ("Current home address", "Residential address, including city and country.", 42),
    ("Home-country address", "Permanent home-country address, if different.", 42),
    ("Home-country phone", "Include country code. A family contact number can be recorded if appropriate.", 25),
], required=("Work email",), dates=("Date of birth",), dropdowns={"Gender":"Gender", "Marital status":"Marital"})

sheet("Emergency Contacts", "04 / EMERGENCY CONTACTS", "One row per contact. Repeat the employee's work email for additional contacts. These are private contacts, not the colleague directory.", [
    EMAIL, ("Contact name", "Full name of emergency contact.", 29),
    ("Relationship", "For example spouse, parent, sibling or friend.", 25),
    ("Contact phone", "Include country code.", 25),
    ("Contact email", "Optional; leave blank if unavailable.", 30),
], required=("Work email", "Contact name", "Relationship", "Contact phone"), rows=60)

sheet("Dependants", "05 / FAMILY & DEPENDANTS", "One row per dependant. Use a simple family reference such as D1 or D2 within each employee's family. Use the same reference in Documents.", [
    EMAIL, ("Family reference", "D1, D2, etc. Unique within this employee's family; not a government ID.", 20),
    ("Dependant full name", "Legal name as shown on passport or ID.", 30),
    ("Relationship", "Spouse, Child, Parent or Other. Explain Other in notes.", 20),
    ("Date of birth", "Required when adding a dependant in the app.", 20),
    ("Nationality", "Nationality of the dependant.", 23),
    ("Contact phone", "For a child use the parent/guardian's number. Include country code.", 25),
    ("Email", "Optional; leave blank for a child without an email.", 28),
    ("Visa required?", "Yes, No or Not confirmed. HR checks eligibility; do not guess.", 22),
    ("Notes for HR", "Clarifications only. Record passport/ID and visa files in Documents.", 38),
], required=("Work email", "Family reference", "Dependant full name", "Relationship", "Date of birth", "Nationality", "Contact phone", "Visa required?"), dropdowns={"Relationship":"Relationship", "Visa required?":"Visa"}, dates=("Date of birth",), rows=100)

sheet("Bank Details", "06 / BANK DETAILS — CONFIDENTIAL", "Complete only if authorised. Share this workbook through an approved private channel. Account numbers are text so leading zeros stay intact.", [
    EMAIL, NAME, ("Account holder name", "Name on the bank account.", 30),
    ("Bank name", "Official bank name.", 28),
    ("Account number", "Type exactly, including leading zeros. Do not use a formula.", 28),
    ("IBAN", "Full IBAN, if applicable. Never enter a card number, PIN or password.", 35),
    ("SWIFT code", "Bank SWIFT/BIC, if applicable.", 23),
    ("Branch", "Branch name, if applicable.", 28),
], required=("Work email",))

sheet("Pay Details", "07 / PAY DETAILS — AUTHORISED HR ONLY", "Optional confidential section. An executive assistant need not complete this. HR must check figures; filling this sheet does not approve payroll.", [
    EMAIL, NAME, ("Base monthly salary", "Amount only; no currency symbol. Use Currency for the currency code.", 25),
    ("Currency", "Approved currency code, e.g. OMR. Do not assume OMR for every country.", 19),
    ("Housing allowance", "Monthly amount; blank means not supplied, not zero.", 24),
    ("Transport allowance", "Monthly amount; blank means not supplied, not zero.", 24),
    ("Pay frequency", "Choose the applicable pay frequency.", 22),
    ("Payment method", "Choose Bank Transfer, Cheque or Cash.", 23),
    ("Social insurance number", "PASI/GOSI or equivalent, if applicable. Keep as text.", 30),
], required=("Work email",), dropdowns={"Pay frequency":"Frequency", "Payment method":"Payment"}, numeric=("Base monthly salary", "Housing allowance", "Transport allowance"))
sheet("Other Allowances", "08 / OTHER ALLOWANCES — AUTHORISED HR ONLY", "One row per additional allowance. Do not repeat housing or transport already entered in Pay Details.", [
    EMAIL, ("Allowance name", "Name of additional allowance.", 30), ("Monthly amount", "Amount only, in the currency on Pay Details.", 24),
], required=("Work email", "Allowance name", "Monthly amount"), numeric=("Monthly amount",), rows=60)

sheet("Documents", "09 / DOCUMENT & PHOTO CHECKLIST", "One row per document or photo. Attach the actual files separately; a filename is not an upload. HR decides which requirements apply and verifies documents in the app.", [
    EMAIL, ("Document reference", "A unique label within this employee, e.g. DOC1. Also used on Extra Document Fields.", 23),
    ("Document for", "Employee or Dependant.", 21),
    ("Family reference", "Use the D1/D2 reference from Dependants only for family documents.", 21),
    ("Document type", "Use the dropdown. A dependant needs passport or ID, plus visa when required. Photo is a profile attachment, not a verification document.", 31),
    ("Requirement name in app", "Copy the exact requirement name if HR added a custom document. Otherwise use the same label as Document type.", 31),
    ("File name", "Exact separate filename with extension. Do not paste public cloud links or embed files in cells.", 38),
    ("Availability", "Already in app: do not upload another copy. Not available yet: HR will follow up.", 26),
    ("Document / card number", "Identity number or insurance member number. Preserve all zeros and letters.", 29),
    ("Issuing authority / provider", "Government authority or insurance provider, as applicable.", 31),
    ("Issuing country", "Country that issued the document, if applicable.", 23),
    ("Issue date", "Leave blank if this document has no issue date.", 20),
    ("Expiry date", "Leave blank if it does not expire. Never invent an expiry date.", 20),
    ("Notes for HR", "Visa/work-permit details are completed by HR. HR verifies files in the app; workbook entry is not approval.", 42),
], required=("Work email", "Document reference", "Document for", "Document type", "Availability"), dropdowns={"Document for":"DocumentFor", "Document type":"DocumentType", "Availability":"Availability"}, dates=("Issue date", "Expiry date"), rows=100)
sheet("Extra Document Fields", "10 / DOCUMENT-SPECIFIC DETAILS", "Use this for education and any extra fields HR configured. For an education degree add three rows: Degree / qualification, Institution, Graduation year.", [
    EMAIL, ("Document reference", "Match Documents, e.g. DOC1.", 24),
    ("Field label in app", "Exact label, such as Institution or Graduation year. Do not invent a field if unsure; ask HR.", 34),
    ("Answer", "Value for that field. Use YYYY-MM-DD for date answers and four digits for years.", 46),
], required=("Work email", "Document reference", "Field label in app", "Answer"), rows=100)

base(start, "VIA / EXECUTIVE PROFILE PACK", "A clear collection template for the CEO, Country Managers and other colleagues who need help completing their profile.", 5)
start.sheet_properties.tabColor = TEAL
for column, width in zip("ABCDE", [9, 28, 30, 30, 30]): start.column_dimensions[column].width = width
start.freeze_panes = "A7"
instructions = [
    ("START", "How to use this workbook", "Fill one row per employee on People, Employment, Personal and (if authorised) Bank Details / Pay Details. Add extra rows on the contact, family and document sheets as needed."),
    ("01", "Use the employee's work email", "The same work email must be used on every sheet. Copy existing employee numbers exactly. Do not use a name alone to match records and do not use the assistant's email."),
    ("02", "The CEO has no supervisor", "In Employment choose No supervisor - CEO. Leave supervisor email and employee number blank. For a Country Manager, enter the actual supervisor. No supervisor does not mean automatic self-approval."),
    ("03", "Who completes what?", "An executive assistant may collect personal details, contacts and files. HR completes and checks employment, visa/work-permit details and confidential pay. Do not give an assistant HR access merely to fill this file."),
    ("04", "Do not guess", "Leave unknown details blank and explain them in People → Notes for HR. A blank cell means not supplied: it must not erase an existing value. A zero is a real value, not a placeholder."),
    ("05", "Use existing setup names", "Department, position, location, employment type, grade, project and cost centre should use their exact app names. Do not create new setup values merely because a spelling differs."),
    ("06", "Dates and numbers", "Use real dates such as 07-Oct-2026. Phone, account and document numbers are text to preserve leading zeros. Salary amounts are numbers without currency symbols. Do not enter passwords, PINs or login codes."),
    ("07", "Documents and photo", "List each document on Documents and supply its actual file separately, with the same filename. Do not embed images into Excel. Employee documents support PDF/JPG/PNG up to 10 MB; provide profile photos as a separate JPG or PNG for review."),
    ("08", "Family documents", "Record each spouse/child in Dependants. Use D1, D2 etc. to link their passport or ID and visa in Documents. A child's contact phone may be the parent/guardian's phone."),
    ("09", "Health insurance", "List the employee's own insurance card. Shared Tables of Benefits are uploaded once by HR and assigned to the appropriate employees or scheme; do not duplicate the shared file for every profile."),
    ("10", "Review before sending", "HR checks People → HR review and enters the reviewer. Amber cells indicate missing key details; red supervisor cells indicate a CEO conflict. Spreadsheet checks are reminders, not final validation."),
    ("11", "Return for import review", "Upload the completed workbook back to this conversation with the separate documents. We must preview exact employee matches, proposed changes and missing information before updating records. Existing employees must be updated, not created again."),
    ("IMPORTANT", "Not a one-click app import", "The current app's bulk importer handles basic new-employee records only. It does not apply this entire workbook or update existing complete profiles. Do not upload all these sheets to that wizard; unsupported information could be missed."),
    ("HR NOTE", "CEO setup before import", "The CEO no-supervisor exception has been implemented in code; confirm it is deployed before import. Do not assign a dummy supervisor. An existing incorrect reporting line must be explicitly cleared by HR. Request approval permissions remain separate and unchanged."),
    ("PRIVATE", "Keep this workbook confidential", "This file may contain family, banking and pay information when filled. Share only with authorised people through an approved private channel. An assistant should not receive completed pay/bank sheets unless authorised. Do not add the completed file to a public repository."),
]
start.cell(6, 1, "STEP")
start.cell(6, 2, "GUIDE")
start.merge_cells("C6:E6")
start.cell(6, 3, "WHAT TO DO")
for c in start[6]:
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.font = Font(name="Calibri", bold=True, color="FFFFFF")
start.row_dimensions[6].height = 26
for row, (step, title, text) in enumerate(instructions, 7):
    start.cell(row, 1, step)
    start.cell(row, 2, title)
    start.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
    start.cell(row, 3, text)
    start.row_dimensions[row].height = 65
    for c in start[row]:
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.font = Font(name="Calibri", size=11, color=INK, bold=c.column < 3)
        c.fill = PatternFill("solid", fgColor=PALE if row % 2 else "FFFFFF")
start.print_area = f"A1:E{6+len(instructions)}"
start.sheet_view.zoomScale = 85

# Put the guide first, the private choices last; preserve worksheet names as import anchors.
wb._sheets = [start] + [ws for ws in wb.worksheets if ws not in (start, lists)] + [lists]
wb.active = 0
wb.properties.title = "VIA Executive Profile Collection Template"
wb.properties.subject = "HR-reviewed collection; no automatic approval or import"
wb.properties.creator = "VIA HR Application"
wb.properties.description = "Blank template, version 1.0. Match by verified work email; preserve existing data; CEO has no supervisor."
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)

# Reopen the actual artifact to validate its structure, empty input rows and workbook references.
check = load_workbook(OUTPUT)
assert len(check.sheetnames) == 12
assert check["Choices"].sheet_state == "hidden"
assert check.active.title == "Start Here"
for ws in check:
    if ws.title in ("Start Here", "Choices"): continue
    assert ws["A6"].value == "Work email"
    assert ws.freeze_panes == "C7"
    assert len(ws.tables) == 1
    assert all(cell.value is None for row in ws.iter_rows(min_row=7) for cell in row)
    for dv in ws.data_validations.dataValidation:
        if dv.type == "list": assert dv.formula1[1:] in check.defined_names
assert check["Bank Details"]["E7"].number_format == "@"
assert check["Employment"]["I7"].number_format == "dd-mmm-yyyy"
assert check["Documents"].max_row == 106
print(f"Created and verified: {OUTPUT}")
print(f"Worksheets: {len(check.sheetnames) - 1} visible; dropdowns, date checks and blank entry rows verified.")
