import base64
import io
import math
import os
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path
from threading import Lock
from typing import Any

import fitz
import pytesseract
from docx import Document
from fastapi import FastAPI, HTTPException
from fastembed import TextEmbedding
from PIL import Image
from pydantic import BaseModel, Field

app = FastAPI(docs_url=None, redoc_url=None, openapi_url=None)

MAX_BYTES = 10 * 1024 * 1024
Image.MAX_IMAGE_PIXELS = 25_000_000
MODEL_NAME = os.getenv(
    "EMBEDDING_MODEL", "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
)
MODEL_CACHE_DIR = os.getenv("EMBEDDING_CACHE_DIR", "/models")
embedding_model = TextEmbedding(model_name=MODEL_NAME, cache_dir=MODEL_CACHE_DIR, threads=4)
embedding_lock = Lock()
SKILLS = {
    "accounting", "autocad", "business development", "customs clearance", "data analysis",
    "excel", "finance", "freight forwarding", "google workspace", "health and safety",
    "human resources", "inventory management", "leadership", "logistics", "operations",
    "payroll", "power bi", "procurement", "project management", "python", "quality assurance",
    "sales", "sap", "supply chain", "transportation", "warehousing", "risk management",
}
LANGUAGES = {"english", "arabic", "french", "hindi", "urdu", "spanish", "german", "mandarin"}


class ExtractRequest(BaseModel):
    fileName: str = Field(min_length=1, max_length=220)
    mimeType: str = Field(default="application/octet-stream", max_length=150)
    contentBase64: str = Field(min_length=1, max_length=14_000_000)


class SimilarityRequest(BaseModel):
    vacancyText: str = Field(min_length=1, max_length=20_000)
    candidateText: str = Field(min_length=1, max_length=30_000)


class SimilarityBatchCandidate(BaseModel):
    candidateId: str = Field(min_length=1, max_length=100)
    candidateText: str = Field(min_length=1, max_length=30_000)


class SimilaritiesRequest(BaseModel):
    vacancyText: str = Field(min_length=1, max_length=20_000)
    candidates: list[SimilarityBatchCandidate] = Field(min_length=1, max_length=500)


def clean_text(value: str) -> str:
    value = value.replace("\x00", " ")
    return re.sub(r"[ \t]+", " ", re.sub(r"\r\n?", "\n", value)).strip()


def quality(text: str) -> float:
    if not text:
        return 0.0
    printable = sum(char.isprintable() or char in "\n\t" for char in text) / len(text)
    words = re.findall(r"[A-Za-z]{2,}", text)
    density = min(1.0, len(words) / 180.0)
    return round(printable * 0.45 + density * 0.55, 3)


def pdf_text(data: bytes) -> tuple[str, bool]:
    document = fitz.open(stream=data, filetype="pdf")
    if len(document) > 100:
        raise ValueError("The PDF has too many pages")
    direct = "\n".join(page.get_text("text") for page in document)
    if quality(direct) >= 0.55:
        return direct, False
    pages = []
    for page in list(document)[:20]:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        pages.append(
            pytesseract.image_to_string(Image.open(io.BytesIO(pixmap.tobytes("png"))), lang="eng")
        )
    return "\n".join(pages), True


def docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > 500 or sum(member.file_size for member in members) > 50 * 1024 * 1024:
            raise ValueError("The Word document expands beyond the safe processing limit")
    document = Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + tables)


def doc_text(data: bytes) -> str:
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "cv.doc"
        path.write_bytes(data)
        result = subprocess.run(
            ["antiword", str(path)], capture_output=True, check=False, timeout=30, text=True
        )
        if result.returncode != 0:
            raise ValueError("The legacy Word document could not be read.")
        return result.stdout


def image_text(data: bytes) -> str:
    return pytesseract.image_to_string(Image.open(io.BytesIO(data)), lang="eng")


def unique_matches(pattern: str, text: str) -> list[str]:
    return list(dict.fromkeys(match.strip(" -:•\t") for match in re.findall(pattern, text, re.I | re.M) if match.strip()))


def structured(text: str, file_name: str) -> tuple[dict[str, Any], dict[str, float], list[str]]:
    lower = text.lower()
    fields: dict[str, Any] = {}
    confidence: dict[str, float] = {}
    evidence: list[str] = []
    email = re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", text, re.I)
    phone = re.search(r"(?:\+|00)?\d[\d\s().-]{7,}\d", text)
    years = re.search(r"\b(\d{1,2})\+?\s+years?\s+(?:of\s+)?(?:professional\s+)?experience\b", text, re.I)
    if email:
        fields["email"] = email.group(0).lower(); confidence["email"] = .99
    if phone:
        fields["phone"] = phone.group(0).strip(); confidence["phone"] = .88
    if years:
        fields["yearsOfExperience"] = int(years.group(1)); confidence["yearsOfExperience"] = .78
        evidence.append(years.group(0))
    skills = sorted(skill for skill in SKILLS if re.search(rf"\b{re.escape(skill)}\b", lower))
    if skills:
        fields["skills"] = skills; confidence["skills"] = .82
    languages = sorted(language.title() for language in LANGUAGES if re.search(rf"\b{language}\b", lower))
    if languages:
        fields["languages"] = languages; confidence["languages"] = .72
    certifications = unique_matches(r"^(?:certifications?|licenses?)\s*[:|-]\s*(.+)$", text)
    education = unique_matches(r"^(?:education|qualifications?)\s*[:|-]\s*(.+)$", text)
    if certifications:
        fields["certifications"] = certifications[:20]; confidence["certifications"] = .67
    if education:
        fields["education"] = education[:20]; confidence["education"] = .67
    eligibility = re.search(r"(?:work|visa)\s+(?:eligibility|authorization)\s*[:|-]?\s*([^\n]{2,100})", text, re.I)
    if eligibility:
        fields["workEligibility"] = eligibility.group(1).strip(); confidence["workEligibility"] = .68
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    history: list[dict[str, str]] = []
    date_range = re.compile(
        r"(?i)(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|"
        r"aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?|\d{1,2})?\s*"
        r"(?:19|20)\d{2}\s*(?:-|to)\s*(?:present|current|(?:19|20)\d{2})"
    )
    for index, line in enumerate(lines):
        dates = date_range.search(line)
        if not dates:
            continue
        nearby = lines[max(0, index - 2):index + 1]
        heading = next((item for item in reversed(nearby) if not date_range.search(item)), line)
        parts = re.split(r"\s+(?:at|@|\||,)\s+", heading, maxsplit=1, flags=re.I)
        date_parts = re.split(r"\s*(?:-|to)\s*", dates.group(0), maxsplit=1, flags=re.I)
        entry: dict[str, str] = {
            "title": parts[0][:160],
            "startDate": date_parts[0].strip(),
            "endDate": date_parts[-1].strip(),
        }
        if len(parts) > 1:
            entry["employer"] = parts[1][:160]
        history.append(entry)
    if history:
        fields["employmentHistory"] = history[:30]; confidence["employmentHistory"] = .62
        current = history[0]
        fields.setdefault("currentTitle", current["title"]); confidence.setdefault("currentTitle", .58)
        if current.get("employer"):
            fields.setdefault("currentCompany", current["employer"]); confidence.setdefault("currentCompany", .58)
    first_lines = [line.strip() for line in text.splitlines() if 2 < len(line.strip()) < 80][:8]
    name_line = next((line for line in first_lines if re.fullmatch(r"[A-Za-z][A-Za-z .'-]{3,60}", line) and "curriculum" not in line.lower()), None)
    if name_line:
        parts = name_line.split()
        if len(parts) >= 2:
            fields["firstName"] = parts[0]; fields["lastName"] = " ".join(parts[1:])
            confidence["firstName"] = .72; confidence["lastName"] = .72
    if not fields.get("firstName"):
        stem = re.sub(r"(?i)\b(cv|resume|final|updated|copy)\b|\d+", " ", Path(file_name).stem)
        parts = re.sub(r"[_().-]+", " ", stem).split()
        if len(parts) >= 2:
            fields["firstName"] = parts[0]; fields["lastName"] = " ".join(parts[1:])
            confidence["firstName"] = .4; confidence["lastName"] = .4
    return fields, confidence, evidence[:30]


def ranking_text(text: str, fields: dict[str, Any]) -> str:
    """Remove identity and protected-characteristic text before relevance matching."""
    redacted = text
    for key in ("email", "phone", "firstName", "lastName"):
        value = fields.get(key)
        if isinstance(value, str) and value.strip():
            redacted = re.sub(re.escape(value.strip()), " ", redacted, flags=re.I)
    excluded_heading = re.compile(
        r"(?i)^\s*(?:name|date of birth|dob|age|gender|sex|marital status|religion|nationality|"
        r"ethnicity|photo|photograph)\s*[:|-]"
    )
    lines = [line for line in redacted.splitlines() if not excluded_heading.search(line)]
    return clean_text("\n".join(lines))[:30_000]


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "embeddingModel": MODEL_NAME}


@app.post("/v1/similarity")
def similarity(request: SimilarityRequest) -> dict[str, Any]:
    vacancy = " ".join(request.vacancyText.split())
    candidate = " ".join(request.candidateText.split())
    with embedding_lock:
        vectors = list(embedding_model.embed([vacancy, candidate], batch_size=2))
    left, right = vectors
    denominator = math.sqrt(float(left @ left)) * math.sqrt(float(right @ right))
    cosine = float(left @ right) / denominator if denominator else 0.0
    return {
        "score": round(max(0.0, min(1.0, cosine)) * 100, 2),
        "model": MODEL_NAME,
        "local": True,
    }


@app.post("/v1/similarities")
def similarities(request: SimilaritiesRequest) -> dict[str, Any]:
    """Embed a vacancy and a pool batch once so large Candidate Pool scans stay fast."""
    texts = [" ".join(request.vacancyText.split())] + [
        " ".join(candidate.candidateText.split()) for candidate in request.candidates
    ]
    with embedding_lock:
        vectors = list(embedding_model.embed(texts, batch_size=min(64, len(texts))))
    vacancy_vector = vectors[0]
    vacancy_length = math.sqrt(float(vacancy_vector @ vacancy_vector))
    results = []
    for candidate, vector in zip(request.candidates, vectors[1:]):
        denominator = vacancy_length * math.sqrt(float(vector @ vector))
        cosine = float(vacancy_vector @ vector) / denominator if denominator else 0.0
        results.append({
            "candidateId": candidate.candidateId,
            "score": round(max(0.0, min(1.0, cosine)) * 100, 2),
        })
    return {"model": MODEL_NAME, "local": True, "results": results}


@app.post("/v1/extract")
def extract(request: ExtractRequest) -> dict[str, Any]:
    try:
        data = base64.b64decode(request.contentBase64, validate=True)
    except Exception as exc:
        raise HTTPException(400, "Invalid document encoding") from exc
    if not data or len(data) > MAX_BYTES:
        raise HTTPException(413, "Document is empty or too large")
    extension = Path(request.fileName).suffix.lower()
    warnings: list[str] = []
    try:
        if extension == ".pdf" or data.startswith(b"%PDF"):
            text, used_ocr = pdf_text(data)
            route = "OCR Required" if used_ocr else "Searchable PDF"
        elif extension == ".docx" or data.startswith(b"PK"):
            text, route = docx_text(data), "Word Document"
        elif extension == ".doc":
            text, route = doc_text(data), "Word Document"
        elif request.mimeType.startswith("image/"):
            text, route = image_text(data), "OCR Required"
        elif request.mimeType.startswith("text/"):
            text, route = data.decode("utf-8", errors="replace"), "Direct Text"
        else:
            raise ValueError("Unsupported CV format")
    except Exception as exc:
        raise HTTPException(422, f"The CV could not be read: {str(exc)[:180]}") from exc
    text = clean_text(text)[:1_500_000]
    score = quality(text)
    if score < .35:
        raise HTTPException(422, "The CV is unreadable and requires HR review")
    if score < .6:
        warnings.append("Some CV content was unclear. HR should verify the proposed information.")
    if route == "OCR Required":
        warnings.append("The document was image-based, so OCR was used. Verify low-confidence fields.")
    fields, confidence, evidence = structured(text, request.fileName)
    if not fields.get("email") and not fields.get("phone"):
        warnings.append("No reliable email address or phone number was found in the CV.")
    return {
        "fields": fields,
        "confidence": confidence,
        "warnings": warnings,
        "evidence": evidence,
        "documentRoute": route,
        "textQuality": score,
        "semanticText": ranking_text(text, fields),
    }
